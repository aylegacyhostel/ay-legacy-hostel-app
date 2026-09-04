from flask import Flask, render_template, request, redirect, url_for, flash, session, send_file, abort
import sqlite3, os, io
from pathlib import Path
from functools import wraps
from datetime import datetime
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from database import db, IS_POSTGRES, integrity_error_types

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY','change-this-secret-key')
BASE = Path(__file__).parent
UPLOADS = BASE/'static'/'uploads'
UPLOADS.mkdir(parents=True, exist_ok=True)
ALLOWED_EXTENSIONS = {'png','jpg','jpeg','webp'}

COMPANY = {
    'parent':'AY Legacy Group', 'name':'AY Legacy Hostel',
    'location':'Sunyani – Fiapra, behind Fiapre community park',
    'phone':'0546188278', 'whatsapp':'233546188278', 'email':'aylegacyhostel@gmail.com'
}


def add_column(conn, table, coldef):
    col = coldef.split()[0]
    if IS_POSTGRES:
        exists = conn.execute(
            "SELECT 1 FROM information_schema.columns WHERE table_schema='public' AND table_name=? AND column_name=?",
            (table, col),
        ).fetchone()
        if not exists:
            conn.execute(f'ALTER TABLE {table} ADD COLUMN {coldef}')
    else:
        existing = [r['name'] for r in conn.execute(f'PRAGMA table_info({table})')]
        if col not in existing:
            conn.execute(f'ALTER TABLE {table} ADD COLUMN {coldef}')


def init_db():
    conn = db()
    if IS_POSTGRES:
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS users(
            id SERIAL PRIMARY KEY, username TEXT UNIQUE NOT NULL, password TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'staff', full_name TEXT, active INTEGER DEFAULT 1);
        CREATE TABLE IF NOT EXISTS rooms(
            id SERIAL PRIMARY KEY, room_number TEXT UNIQUE NOT NULL, capacity INTEGER NOT NULL DEFAULT 1,
            floor TEXT, room_type TEXT, monthly_rate DOUBLE PRECISION DEFAULT 0, status TEXT DEFAULT 'Available', notes TEXT);
        CREATE TABLE IF NOT EXISTS enquiries(
            id SERIAL PRIMARY KEY, name TEXT NOT NULL, phone TEXT NOT NULL, email TEXT, school TEXT, program TEXT,
            year_level TEXT, preferred_room TEXT, message TEXT, status TEXT DEFAULT 'New', created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS students(
            id SERIAL PRIMARY KEY, student_id TEXT UNIQUE, first_name TEXT NOT NULL, last_name TEXT NOT NULL,
            gender TEXT, phone TEXT, email TEXT, school TEXT, program TEXT, year_level TEXT, date_of_birth TEXT, hometown TEXT,
            room_number TEXT, bed_space TEXT, check_in_date TEXT, expected_check_out TEXT, status TEXT DEFAULT 'Active',
            parent_name TEXT, parent_phone TEXT, parent_email TEXT, parent_relationship TEXT,
            emergency_name TEXT, emergency_phone TEXT, emergency_relationship TEXT, medical_notes TEXT,
            id_type TEXT, id_number TEXT, notes TEXT, photo TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS payments(
            id SERIAL PRIMARY KEY, student_id INTEGER NOT NULL REFERENCES students(id), academic_year TEXT,
            payment_type TEXT DEFAULT 'Hostel Fee', amount_due DOUBLE PRECISION NOT NULL DEFAULT 0,
            amount_paid DOUBLE PRECISION NOT NULL DEFAULT 0, payment_date TEXT, payment_method TEXT,
            reference TEXT, notes TEXT, receipt_no TEXT, created_by TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS audit_logs(
            id SERIAL PRIMARY KEY, username TEXT, action TEXT NOT NULL, entity_type TEXT NOT NULL,
            entity_id INTEGER, details TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP);
        """)
    else:
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE NOT NULL, password TEXT NOT NULL, role TEXT NOT NULL DEFAULT 'staff', full_name TEXT, active INTEGER DEFAULT 1);
        CREATE TABLE IF NOT EXISTS rooms(id INTEGER PRIMARY KEY AUTOINCREMENT, room_number TEXT UNIQUE NOT NULL, capacity INTEGER NOT NULL DEFAULT 1, floor TEXT, room_type TEXT, monthly_rate REAL DEFAULT 0, status TEXT DEFAULT 'Available', notes TEXT);
        CREATE TABLE IF NOT EXISTS enquiries(id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, phone TEXT NOT NULL, email TEXT, school TEXT, program TEXT, year_level TEXT, preferred_room TEXT, message TEXT, status TEXT DEFAULT 'New', created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS students(
            id INTEGER PRIMARY KEY AUTOINCREMENT, student_id TEXT UNIQUE, first_name TEXT NOT NULL, last_name TEXT NOT NULL,
            gender TEXT, phone TEXT, email TEXT, school TEXT, program TEXT, year_level TEXT, date_of_birth TEXT, hometown TEXT,
            room_number TEXT, bed_space TEXT, check_in_date TEXT, expected_check_out TEXT, status TEXT DEFAULT 'Active',
            parent_name TEXT, parent_phone TEXT, parent_email TEXT, parent_relationship TEXT,
            emergency_name TEXT, emergency_phone TEXT, emergency_relationship TEXT, medical_notes TEXT,
            id_type TEXT, id_number TEXT, notes TEXT, photo TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS payments(
            id INTEGER PRIMARY KEY AUTOINCREMENT, student_id INTEGER NOT NULL, academic_year TEXT, payment_type TEXT DEFAULT 'Hostel Fee',
            amount_due REAL NOT NULL DEFAULT 0, amount_paid REAL NOT NULL DEFAULT 0, payment_date TEXT, payment_method TEXT,
            reference TEXT, notes TEXT, receipt_no TEXT, created_by TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(student_id) REFERENCES students(id));
        CREATE TABLE IF NOT EXISTS audit_logs(
            id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT, action TEXT NOT NULL, entity_type TEXT NOT NULL,
            entity_id INTEGER, details TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        """)
    add_column(conn, 'users', 'full_name TEXT')
    add_column(conn, 'users', 'active INTEGER DEFAULT 1')
    add_column(conn, 'students', 'photo TEXT')
    add_column(conn, 'payments', 'receipt_no TEXT')
    add_column(conn, 'payments', 'created_by TEXT')
    if conn.execute('SELECT COUNT(*) c FROM users').fetchone()['c'] == 0:
        conn.execute('INSERT INTO users(username,password,role,full_name) VALUES(?,?,?,?)',
                     ('admin', generate_password_hash('admin123'), 'admin', 'Administrator'))
    else:
        u = conn.execute("SELECT * FROM users WHERE username='admin'").fetchone()
        if u and u['password'] == 'admin123':
            conn.execute("UPDATE users SET password=?, role='admin' WHERE username='admin'",
                         (generate_password_hash('admin123'),))
    conn.commit()
    conn.close()


def audit(conn, action, entity_type, entity_id=None, details=''):
    conn.execute(
        'INSERT INTO audit_logs(username,action,entity_type,entity_id,details) VALUES(?,?,?,?,?)',
        (session.get('user', 'public'), action, entity_type, entity_id, details),
    )

@app.context_processor
def inject_company(): return dict(company=COMPANY)

def login_required(fn):
    @wraps(fn)
    def w(*a,**k):
        if 'user' not in session: return redirect(url_for('login'))
        return fn(*a,**k)
    return w

def admin_required(fn):
    @wraps(fn)
    def w(*a,**k):
        if session.get('role')!='admin': abort(403)
        return fn(*a,**k)
    return w

def allowed_file(name): return '.' in name and name.rsplit('.',1)[1].lower() in ALLOWED_EXTENSIONS

def student_totals(conn,sid):
    return conn.execute('SELECT COALESCE(SUM(amount_due),0) due,COALESCE(SUM(amount_paid),0) paid,COALESCE(SUM(amount_due-amount_paid),0) balance FROM payments WHERE student_id=?',(sid,)).fetchone()

@app.route('/login',methods=['GET','POST'])
def login():
    if request.method=='POST':
        conn=db(); u=conn.execute('SELECT * FROM users WHERE username=? AND active=1',(request.form.get('username','').strip(),)).fetchone(); conn.close()
        if u and (check_password_hash(u['password'],request.form.get('password','')) if u['password']!='admin123' else request.form.get('password')=='admin123'):
            session.update(user=u['username'], role=u['role'], full_name=u['full_name'] or u['username']); return redirect(url_for('dashboard'))
        flash('Invalid username or password.','danger')
    return render_template('login.html')

@app.route('/logout')
def logout(): session.clear(); return redirect(url_for('login'))


@app.route('/account/password', methods=['GET','POST'])
@login_required
def change_password():
    if request.method == 'POST':
        current = request.form.get('current_password','')
        new_password = request.form.get('new_password','')
        confirm = request.form.get('confirm_password','')
        conn = db()
        u = conn.execute('SELECT * FROM users WHERE username=?',(session.get('user'),)).fetchone()
        if not u or not check_password_hash(u['password'], current):
            conn.close(); flash('Current password is incorrect.','danger')
        elif len(new_password) < 8:
            conn.close(); flash('New password must be at least 8 characters.','danger')
        elif new_password != confirm:
            conn.close(); flash('New passwords do not match.','danger')
        else:
            conn.execute('UPDATE users SET password=? WHERE id=?',(generate_password_hash(new_password),u['id']))
            audit(conn,'change_password','user',u['id'])
            conn.commit(); conn.close(); flash('Password changed successfully.','success')
            return redirect(url_for('dashboard'))
    return render_template('change_password.html')

@app.route('/audit-logs')
@login_required
@admin_required
def audit_logs():
    conn=db()
    rows=conn.execute('SELECT * FROM audit_logs ORDER BY created_at DESC,id DESC LIMIT 500').fetchall()
    conn.close()
    return render_template('audit_logs.html',rows=rows)

@app.route('/')
@login_required
def dashboard():
    conn=db()
    stats=conn.execute("SELECT COUNT(*) total, SUM(CASE WHEN status='Active' THEN 1 ELSE 0 END) active FROM students").fetchone()
    money=conn.execute('SELECT COALESCE(SUM(amount_due),0) due,COALESCE(SUM(amount_paid),0) paid FROM payments').fetchone()
    beds=conn.execute("SELECT COALESCE(SUM(capacity),0) total FROM rooms WHERE status!='Maintenance'").fetchone()['total']
    occupied=conn.execute("SELECT COUNT(*) c FROM students WHERE status='Active' AND room_number IS NOT NULL AND room_number<>''").fetchone()['c']
    balances=conn.execute('''SELECT s.id,s.first_name||' '||s.last_name student_name,s.room_number,COALESCE(SUM(p.amount_due),0) due,COALESCE(SUM(p.amount_paid),0) paid,COALESCE(SUM(p.amount_due-p.amount_paid),0) balance FROM students s LEFT JOIN payments p ON p.student_id=s.id GROUP BY s.id HAVING COALESCE(SUM(p.amount_due-p.amount_paid),0)>0 ORDER BY balance DESC LIMIT 10''').fetchall()
    recent=conn.execute("SELECT p.*,s.first_name||' '||s.last_name student_name FROM payments p JOIN students s ON s.id=p.student_id ORDER BY COALESCE(NULLIF(p.payment_date,''),CAST(p.created_at AS TEXT)) DESC LIMIT 8").fetchall()
    new_enquiries=conn.execute("SELECT COUNT(*) c FROM enquiries WHERE status='New'").fetchone()['c']; conn.close()
    return render_template('dashboard.html',total_students=stats['active'] or 0,total_due=money['due'],total_paid=money['paid'],outstanding=money['due']-money['paid'],beds=beds,occupied=occupied,new_enquiries=new_enquiries,balances=balances,recent=recent)

@app.route('/students')
@login_required
def students():
    q=request.args.get('q','').strip(); year=request.args.get('year','').strip(); status=request.args.get('status','').strip(); room=request.args.get('room','').strip(); program=request.args.get('program','').strip()
    sql="""SELECT s.*,COALESCE(SUM(p.amount_due),0) total_due,COALESCE(SUM(p.amount_paid),0) total_paid,COALESCE(SUM(p.amount_due-p.amount_paid),0) balance FROM students s LEFT JOIN payments p ON p.student_id=s.id WHERE 1=1"""; params=[]
    if q:
        like=f'%{q}%'; sql+=" AND (s.first_name LIKE ? OR s.last_name LIKE ? OR s.student_id LIKE ? OR s.phone LIKE ? OR s.room_number LIKE ? OR s.program LIKE ?)"; params += [like]*6
    if year: sql+=' AND s.year_level=?'; params.append(year)
    if status: sql+=' AND s.status=?'; params.append(status)
    if room: sql+=' AND s.room_number=?'; params.append(room)
    if program: sql+=' AND s.program=?'; params.append(program)
    sql+=' GROUP BY s.id ORDER BY s.last_name,s.first_name'
    conn=db(); rows=conn.execute(sql,params).fetchall(); rooms=conn.execute('SELECT room_number FROM rooms ORDER BY room_number').fetchall(); programs=conn.execute("SELECT DISTINCT program FROM students WHERE program IS NOT NULL AND TRIM(program)<>'' ORDER BY program").fetchall(); conn.close()
    return render_template('students.html',students=rows,q=q,year=year,status=status,room=room,program=program,rooms=rooms,programs=programs)

def save_student_form(conn,sid=None):
    fields=['student_id','first_name','last_name','gender','phone','email','school','program','year_level','date_of_birth','hometown','room_number','bed_space','check_in_date','expected_check_out','status','parent_name','parent_phone','parent_email','parent_relationship','emergency_name','emergency_phone','emergency_relationship','medical_notes','id_type','id_number','notes']
    vals=[request.form.get(f,'').strip() for f in fields]
    photo=None
    f=request.files.get('photo')
    if f and f.filename and allowed_file(f.filename):
        photo=f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{secure_filename(f.filename)}"; f.save(UPLOADS/photo)
    if sid:
        pairs=[f'{x}=?' for x in fields];
        if photo: pairs.append('photo=?'); vals.append(photo)
        conn.execute(f"UPDATE students SET {','.join(pairs)} WHERE id=?",vals+[sid])
    else:
        if photo: fields.append('photo'); vals.append(photo)
        conn.execute(f"INSERT INTO students({','.join(fields)}) VALUES({','.join(['?']*len(fields))})",vals)

@app.route('/students/new',methods=['GET','POST'])
@login_required
def new_student():
    conn=db(); rooms=conn.execute('SELECT * FROM rooms ORDER BY room_number').fetchall()
    if request.method=='POST':
        try: save_student_form(conn); audit(conn,'create','student',details=request.form.get('student_id','')); conn.commit(); flash('Student added successfully.','success'); conn.close(); return redirect(url_for('students'))
        except integrity_error_types(): flash('Student ID already exists.','danger')
    conn.close(); return render_template('student_form.html',student=None,rooms=rooms)

@app.route('/students/<int:sid>/edit',methods=['GET','POST'])
@login_required
def edit_student(sid):
    conn=db(); student=conn.execute('SELECT * FROM students WHERE id=?',(sid,)).fetchone(); rooms=conn.execute('SELECT * FROM rooms ORDER BY room_number').fetchall()
    if not student: conn.close(); abort(404)
    if request.method=='POST':
        try: save_student_form(conn,sid); audit(conn,'update','student',sid,request.form.get('student_id','')); conn.commit(); flash('Student updated successfully.','success'); conn.close(); return redirect(url_for('student_detail',sid=sid))
        except integrity_error_types(): flash('Student ID already exists.','danger')
    conn.close(); return render_template('student_form.html',student=student,rooms=rooms)

@app.route('/students/<int:sid>')
@login_required
def student_detail(sid):
    conn=db(); s=conn.execute('SELECT * FROM students WHERE id=?',(sid,)).fetchone();
    if not s: conn.close(); abort(404)
    payments=conn.execute('SELECT * FROM payments WHERE student_id=? ORDER BY payment_date DESC,id DESC',(sid,)).fetchall(); totals=student_totals(conn,sid); conn.close()
    return render_template('student_detail.html',student=s,payments=payments,totals=totals)

@app.route('/students/<int:sid>/payment/new',methods=['GET','POST'])
@login_required
def new_payment(sid):
    conn=db(); s=conn.execute('SELECT * FROM students WHERE id=?',(sid,)).fetchone();
    if not s: conn.close(); abort(404)
    if request.method=='POST':
        receipt='AYH-'+datetime.now().strftime('%Y%m%d%H%M%S')
        conn.execute('''INSERT INTO payments(student_id,academic_year,payment_type,amount_due,amount_paid,payment_date,payment_method,reference,notes,receipt_no,created_by) VALUES(?,?,?,?,?,?,?,?,?,?,?)''',(sid,request.form.get('academic_year','').strip(),request.form.get('payment_type','Hostel Fee').strip(),float(request.form.get('amount_due') or 0),float(request.form.get('amount_paid') or 0),request.form.get('payment_date','').strip(),request.form.get('payment_method','').strip(),request.form.get('reference','').strip(),request.form.get('notes','').strip(),receipt,session.get('user')))
        pid=conn.execute('SELECT last_insert_rowid() id').fetchone()['id']; audit(conn,'create','payment',pid,f"student_id={sid}; amount_paid={request.form.get('amount_paid','0')}"); conn.commit(); conn.close(); flash('Payment recorded.','success'); return redirect(url_for('receipt',pid=pid))
    conn.close(); return render_template('payment_form.html',student=s,payment=None)

@app.route('/payments/<int:pid>/edit',methods=['GET','POST'])
@login_required
def edit_payment(pid):
    conn=db(); p=conn.execute('SELECT * FROM payments WHERE id=?',(pid,)).fetchone();
    if not p: conn.close(); abort(404)
    s=conn.execute('SELECT * FROM students WHERE id=?',(p['student_id'],)).fetchone()
    if request.method=='POST':
        conn.execute('''UPDATE payments SET academic_year=?,payment_type=?,amount_due=?,amount_paid=?,payment_date=?,payment_method=?,reference=?,notes=? WHERE id=?''',(request.form.get('academic_year','').strip(),request.form.get('payment_type','Hostel Fee').strip(),float(request.form.get('amount_due') or 0),float(request.form.get('amount_paid') or 0),request.form.get('payment_date','').strip(),request.form.get('payment_method','').strip(),request.form.get('reference','').strip(),request.form.get('notes','').strip(),pid)); audit(conn,'update','payment',pid,f"amount_paid={request.form.get('amount_paid','0')}"); conn.commit(); conn.close(); flash('Payment updated.','success'); return redirect(url_for('student_detail',sid=s['id']))
    conn.close(); return render_template('payment_form.html',student=s,payment=p)

@app.route('/payments')
@login_required
def payments():
    conn=db(); rows=conn.execute("SELECT p.*,s.first_name||' '||s.last_name student_name,s.room_number FROM payments p JOIN students s ON s.id=p.student_id ORDER BY p.payment_date DESC,p.id DESC").fetchall(); conn.close(); return render_template('payments.html',payments=rows)

@app.route('/receipt/<int:pid>')
@login_required
def receipt(pid):
    conn=db(); row=conn.execute("SELECT p.*,s.first_name||' '||s.last_name student_name,s.student_id student_code,s.room_number FROM payments p JOIN students s ON s.id=p.student_id WHERE p.id=?",(pid,)).fetchone(); totals=student_totals(conn,row['student_id']) if row else None; conn.close()
    if not row: abort(404)
    return render_template('receipt.html',p=row,totals=totals)

@app.route('/receipt/<int:pid>/pdf')
@login_required
def receipt_pdf(pid):
    conn=db(); p=conn.execute("SELECT p.*,s.first_name||' '||s.last_name student_name,s.student_id student_code,s.room_number FROM payments p JOIN students s ON s.id=p.student_id WHERE p.id=?",(pid,)).fetchone(); totals=student_totals(conn,p['student_id']) if p else None; conn.close()
    if not p: abort(404)
    bio=io.BytesIO(); doc=SimpleDocTemplate(bio,pagesize=A4,rightMargin=36,leftMargin=36,topMargin=36,bottomMargin=36); styles=getSampleStyleSheet()
    story=[Paragraph('<b>AY LEGACY HOSTEL - PAYMENT RECEIPT</b>',styles['Title']),Paragraph(COMPANY['location'],styles['Normal']),Paragraph('Tel: '+COMPANY['phone'],styles['Normal']),Spacer(1,18)]
    data=[['Receipt No.',p['receipt_no'] or str(p['id'])],['Student',p['student_name']],['Student ID',p['student_code'] or '-'],['Room',p['room_number'] or '-'],['Payment Date',p['payment_date'] or '-'],['Academic Year',p['academic_year'] or '-'],['Payment Type',p['payment_type']],['Amount Due',f"GHS {p['amount_due']:.2f}"],['Amount Paid',f"GHS {p['amount_paid']:.2f}"],['Current Total Balance',f"GHS {totals['balance']:.2f}"],['Method',p['payment_method'] or '-'],['Reference',p['reference'] or '-']]
    t=Table(data,colWidths=[130,330]); t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),0.5,colors.grey),('BACKGROUND',(0,0),(0,-1),colors.whitesmoke),('FONTNAME',(0,0),(0,-1),'Helvetica-Bold'),('PADDING',(0,0),(-1,-1),8)])); story += [t,Spacer(1,18),Paragraph('Thank you for your payment.',styles['Normal'])]; doc.build(story); bio.seek(0)
    return send_file(bio,as_attachment=True,download_name=f"receipt_{p['receipt_no'] or pid}.pdf",mimetype='application/pdf')

@app.route('/arrears')
@login_required
def arrears():
    conn=db(); rows=conn.execute('''SELECT s.id,s.first_name||' '||s.last_name student_name,s.phone,s.parent_phone,s.room_number,s.program,s.year_level,COALESCE(SUM(p.amount_due),0) due,COALESCE(SUM(p.amount_paid),0) paid,COALESCE(SUM(p.amount_due-p.amount_paid),0) balance FROM students s LEFT JOIN payments p ON p.student_id=s.id GROUP BY s.id HAVING COALESCE(SUM(p.amount_due-p.amount_paid),0)>0 ORDER BY balance DESC''').fetchall(); conn.close(); return render_template('arrears.html',rows=rows)

@app.route('/rooms')
@login_required
def rooms():
    conn=db(); rows=conn.execute('''SELECT r.*,COUNT(s.id) occupied FROM rooms r LEFT JOIN students s ON s.room_number=r.room_number AND s.status='Active' GROUP BY r.id ORDER BY r.room_number''').fetchall(); conn.close(); return render_template('rooms.html',rooms=rows)

@app.route('/rooms/new',methods=['GET','POST'])
@login_required
@admin_required
def new_room():
    if request.method=='POST':
        conn=db()
        try:
            conn.execute('INSERT INTO rooms(room_number,capacity,floor,room_type,monthly_rate,status,notes) VALUES(?,?,?,?,?,?,?)',(request.form.get('room_number','').strip(),int(request.form.get('capacity') or 1),request.form.get('floor','').strip(),request.form.get('room_type','').strip(),float(request.form.get('monthly_rate') or 0),request.form.get('status','Available'),request.form.get('notes','').strip())); audit(conn,'create','room',details=request.form.get('room_number','')); conn.commit(); flash('Room added.','success'); return redirect(url_for('rooms'))
        except integrity_error_types(): flash('Room number already exists.','danger')
        finally: conn.close()
    return render_template('room_form.html',room=None)

@app.route('/rooms/<int:rid>/edit',methods=['GET','POST'])
@login_required
@admin_required
def edit_room(rid):
    conn=db(); r=conn.execute('SELECT * FROM rooms WHERE id=?',(rid,)).fetchone()
    if not r: conn.close(); abort(404)
    if request.method=='POST':
        conn.execute('UPDATE rooms SET room_number=?,capacity=?,floor=?,room_type=?,monthly_rate=?,status=?,notes=? WHERE id=?',(request.form.get('room_number','').strip(),int(request.form.get('capacity') or 1),request.form.get('floor','').strip(),request.form.get('room_type','').strip(),float(request.form.get('monthly_rate') or 0),request.form.get('status','Available'),request.form.get('notes','').strip(),rid)); audit(conn,'update','room',rid,request.form.get('room_number','')); conn.commit(); conn.close(); flash('Room updated.','success'); return redirect(url_for('rooms'))
    conn.close(); return render_template('room_form.html',room=r)

@app.route('/users')
@login_required
@admin_required
def users():
    conn=db(); rows=conn.execute('SELECT id,username,role,full_name,active FROM users ORDER BY username').fetchall(); conn.close(); return render_template('users.html',users=rows)

@app.route('/users/new',methods=['GET','POST'])
@login_required
@admin_required
def new_user():
    if request.method=='POST':
        conn=db()
        try:
            conn.execute('INSERT INTO users(username,password,role,full_name,active) VALUES(?,?,?,?,1)',(request.form.get('username','').strip(),generate_password_hash(request.form.get('password','')),request.form.get('role','staff'),request.form.get('full_name','').strip())); audit(conn,'create','user',details=request.form.get('username','')); conn.commit(); flash('Staff account created.','success'); return redirect(url_for('users'))
        except integrity_error_types(): flash('Username already exists.','danger')
        finally: conn.close()
    return render_template('user_form.html')

@app.route('/users/<int:uid>/toggle')
@login_required
@admin_required
def toggle_user(uid):
    if uid==1: flash('Main administrator account cannot be disabled.','danger'); return redirect(url_for('users'))
    conn=db(); conn.execute('UPDATE users SET active=CASE active WHEN 1 THEN 0 ELSE 1 END WHERE id=?',(uid,)); audit(conn,'toggle','user',uid); conn.commit(); conn.close(); return redirect(url_for('users'))

@app.route('/enquiries')
@login_required
def enquiries():
    conn=db(); rows=conn.execute('SELECT * FROM enquiries ORDER BY created_at DESC').fetchall(); conn.close(); return render_template('enquiries.html',enquiries=rows)

@app.route('/enquiries/<int:eid>/status',methods=['POST'])
@login_required
def enquiry_status(eid):
    conn=db(); conn.execute('UPDATE enquiries SET status=? WHERE id=?',(request.form.get('status','New'),eid)); audit(conn,'update_status','enquiry',eid,request.form.get('status','New')); conn.commit(); conn.close(); return redirect(url_for('enquiries'))

@app.route('/apply',methods=['GET','POST'])
def apply():
    if request.method=='POST':
        conn=db(); conn.execute('INSERT INTO enquiries(name,phone,email,school,program,year_level,preferred_room,message) VALUES(?,?,?,?,?,?,?,?)',(request.form.get('name','').strip(),request.form.get('phone','').strip(),request.form.get('email','').strip(),request.form.get('school','').strip(),request.form.get('program','').strip(),request.form.get('year_level','').strip(),request.form.get('preferred_room','').strip(),request.form.get('message','').strip())); conn.commit(); conn.close(); flash('Your enquiry has been submitted. AY Legacy Hostel will contact you.','success'); return redirect(url_for('apply'))
    return render_template('apply.html')

@app.route('/reports/students.xlsx')
@login_required
def students_xlsx():
    conn=db(); rows=conn.execute("SELECT student_id,first_name,last_name,gender,phone,email,school,program,year_level,room_number,bed_space,status,parent_name,parent_phone,emergency_name,emergency_phone FROM students ORDER BY last_name,first_name").fetchall(); conn.close()
    wb=Workbook(); ws=wb.active; ws.title='Students'; headers=list(rows[0].keys()) if rows else ['student_id','first_name','last_name','gender','phone','email','school','program','year_level','room_number','bed_space','status','parent_name','parent_phone','emergency_name','emergency_phone']; ws.append(headers)
    for r in rows: ws.append([r[h] for h in headers])
    for c in ws[1]: c.font=c.font.copy(bold=True)
    bio=io.BytesIO(); wb.save(bio); bio.seek(0); return send_file(bio,as_attachment=True,download_name='AY_Legacy_Hostel_Students.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/reports/arrears.xlsx')
@login_required
def arrears_xlsx():
    conn=db(); rows=conn.execute('''SELECT s.student_id,s.first_name||' '||s.last_name student,s.phone,s.parent_phone,s.room_number,s.program,s.year_level,COALESCE(SUM(p.amount_due),0) due,COALESCE(SUM(p.amount_paid),0) paid,COALESCE(SUM(p.amount_due-p.amount_paid),0) balance FROM students s LEFT JOIN payments p ON p.student_id=s.id GROUP BY s.id HAVING COALESCE(SUM(p.amount_due-p.amount_paid),0)>0 ORDER BY balance DESC''').fetchall(); conn.close()
    wb=Workbook(); ws=wb.active; ws.title='Arrears'; headers=['student_id','student','phone','parent_phone','room_number','program','year_level','due','paid','balance']; ws.append(headers)
    for r in rows: ws.append([r[h] for h in headers])
    for c in ws[1]: c.font=c.font.copy(bold=True)
    bio=io.BytesIO(); wb.save(bio); bio.seek(0); return send_file(bio,as_attachment=True,download_name='AY_Legacy_Hostel_Arrears.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

init_db()

if __name__=='__main__':
    app.run(debug=True,host='0.0.0.0',port=int(os.environ.get('PORT','5000')))
